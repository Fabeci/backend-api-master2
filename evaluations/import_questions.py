# evaluations/import_questions.py
# =============================================================================
# Extraction de questions depuis un fichier (Excel, CSV, Word, PDF, TXT).
#
# Objectif : pré-remplir le formulaire de création d'évaluation / quiz.
# Deux modes :
#   1) STRUCTURÉ (xlsx, csv) : colonnes normalisées (fiable).
#   2) TEXTE (txt, docx, pdf) : heuristiques (questions numérotées, options
#      a) b) c), ligne « Réponse: b »), avec avertissements.
#
# Sortie : { questions: [ {enonce_texte, type_question, points, reponses:[{texte,est_correcte}]} ],
#            warnings: [str], nb: int }
# =============================================================================

import csv
import io
import logging
import re

logger = logging.getLogger(__name__)

TYPE_UNIQUE = "choix_unique"
TYPE_MULTIPLE = "choix_multiple"
TYPE_TEXTE = "texte_long"

_WARN = []


def _add_warning(msg: str) -> None:
    _WARN.append(msg)


# ─────────────────────────────────────────────────────────────────────────────
# 1) STRUCTURÉ — Excel / CSV
# ─────────────────────────────────────────────────────────────────────────────

# Colonnes attendues (tolérant : insensible à la casse/accents)
_COL_ALIASES = {
    "question": "enonce", "enonce": "enonce", "énoncé": "enonce", "intitule": "enonce",
    "type": "type", "points": "points", "point": "points", "bareme": "points",
    "bonne_reponse": "bonne", "réponse": "bonne", "reponse": "bonne", "correcte": "bonne",
    "answer": "bonne",
    "explication": "explication", "indication": "explication",
}
_OPT_PREFIX = "option"


def _norm_header(h: str) -> str:
    h = (h or "").strip().lower()
    for acc, rep in (("é", "e"), ("è", "e"), ("ê", "e"), ("à", "a"), ("î", "i"), ("ô", "o"), ("û", "u"), ("ç", "c")):
        h = h.replace(acc, rep)
    h = re.sub(r"[\s\-]+", "_", h)
    return h


def _letters_to_indices(raw: str, nb_options: int) -> list[int]:
    """'b' -> [1], 'a,c' -> [0,2], '2' -> [1]."""
    out = []
    if not raw:
        return out
    for tok in re.split(r"[;,/|]", str(raw)):
        tok = tok.strip().lower()
        if not tok:
            continue
        if tok.isdigit():
            out.append(int(tok) - 1)
        elif re.fullmatch(r"[a-h]", tok):
            out.append(ord(tok) - ord("a"))
    return [i for i in out if 0 <= i < nb_options]


def _rows_to_questions(rows: list[list]) -> list[dict]:
    if not rows:
        return []
    header = [_norm_header(str(c)) for c in rows[0]]
    idx = {}
    option_cols = []
    for i, h in enumerate(header):
        key = _COL_ALIASES.get(h)
        if key:
            idx.setdefault(key, i)
        elif h.startswith(_OPT_PREFIX) or re.fullmatch(r"o[0-9]+", h) or re.fullmatch(r"(choix|reponse)_?[0-9]+", h):
            option_cols.append(i)

    if "enonce" not in idx:
        _add_warning("Colonne « question » (ou « enonce ») introuvable — vérifiez l'entête.")
        return []

    questions = []
    for r, row in enumerate(rows[1:], start=2):
        def cell(key, default=""):
            i = idx.get(key)
            return (str(row[i]).strip() if i is not None and i < len(row) and row[i] is not None else default)

        enonce = cell("enonce")
        if not enonce:
            continue

        options = []
        for i in option_cols:
            if i < len(row) and row[i] not in (None, ""):
                options.append(str(row[i]).strip())
        # options vides entre-temps → on retire
        options = [o for o in options if o]

        bonne = cell("bonne")
        corrects = _letters_to_indices(bonne, len(options))
        if options and not corrects:
            _add_warning(f"Ligne {r} : aucune bonne réponse reconnue pour « {enonce[:40]}… » (colonne bonne_reponse).")

        type_q = (cell("type") or "").strip().lower()
        if type_q not in (TYPE_UNIQUE, TYPE_MULTIPLE, "texte", "texte_court", "texte_long"):
            type_q = TYPE_MULTIPLE if len(corrects) > 1 else (TYPE_UNIQUE if options else TYPE_TEXTE)
        elif type_q == "texte":
            type_q = TYPE_TEXTE

        try:
            points = float(str(cell("points", "1")).replace(",", "."))
        except Exception:
            points = 1.0

        questions.append({
            "enonce_texte": enonce,
            "type_question": type_q,
            "points": points,
            "indication_reponse": cell("explication"),
            "reponses": [{"texte": o, "est_correcte": (i in corrects)} for i, o in enumerate(options)],
        })
    return questions


def _parse_xlsx(data: bytes) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return _rows_to_questions(rows)


def _parse_csv(data: bytes) -> list[dict]:
    text = data.decode("utf-8-sig", errors="ignore")
    sample = text[:2000]
    delim = ";" if sample.count(";") > sample.count(",") else ","
    rows = [row for row in csv.reader(io.StringIO(text), delimiter=delim)]
    return _rows_to_questions(rows)


# ─────────────────────────────────────────────────────────────────────────────
# 2) TEXTE — Word / PDF / TXT (heuristiques)
# ─────────────────────────────────────────────────────────────────────────────

_Q_START = re.compile(r"^\s*(?:question\s*)?(\d{1,3})\s*[).:\-]\s*(.+)$", re.IGNORECASE)
_OPT = re.compile(r"^\s*\(?([a-hA-H])\s*[)\].:\-]\s*(.+)$")
_ANS = re.compile(r"^\s*(?:bonne\s+)?(?:r[eé]ponse|answer|correct(?:e)?)\s*[:\-]\s*(.+)$", re.IGNORECASE)
_PTS = re.compile(r"\(?\s*(\d+(?:[.,]\d+)?)\s*(?:pts?|points?)\s*\)?", re.IGNORECASE)


def _parse_text(text: str) -> list[dict]:
    lines = [l.rstrip() for l in text.splitlines()]
    questions = []
    cur = None
    for line in lines:
        s = line.strip()
        if not s:
            continue

        m = _ANS.match(s)
        if m and cur is not None:
            cur["_answer"] = m.group(1)
            continue

        mo = _OPT.match(s)
        if mo and cur is not None:
            cur["_opts"].append((mo.group(1).lower(), mo.group(2).strip()))
            continue

        mq = _Q_START.match(s)
        if mq and (len(s) < 200 or s.endswith("?")):
            # nouvelle question
            pts = _PTS.search(s)
            enonce = _PTS.sub("", mq.group(2)).strip()
            cur = {"_enonce": enonce, "_opts": [], "_answer": "", "_points": float(pts.group(1).replace(",", ".")) if pts else 1.0}
            questions.append(cur)
            continue

        # continuation d'énoncé
        if cur is not None and not cur["_opts"]:
            cur["_enonce"] += " " + s

    # convertir
    out = []
    for q in questions:
        opts = q["_opts"]
        corrects = _letters_to_indices(q.get("_answer", ""), len(opts)) if opts else []
        if opts and not corrects:
            _add_warning(f"Aucune bonne réponse détectée pour « {q['_enonce'][:40]}… » (ajoutez « Réponse : b »).")
        reponses = [{"texte": t, "est_correcte": (i in corrects)} for i, (_, t) in enumerate(opts)]
        type_q = TYPE_MULTIPLE if len(corrects) > 1 else (TYPE_UNIQUE if opts else TYPE_TEXTE)
        out.append({
            "enonce_texte": q["_enonce"].strip(),
            "type_question": type_q,
            "points": q["_points"],
            "indication_reponse": "",
            "reponses": reponses,
        })
    return out


def _parse_docx(data: bytes) -> str:
    import docx  # python-docx
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def _parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


# ─────────────────────────────────────────────────────────────────────────────
# API
# ─────────────────────────────────────────────────────────────────────────────

def extract_questions(filename: str, data: bytes) -> dict:
    """Point d'entrée : détecte le format et extrait les questions."""
    _WARN.clear()
    name = (filename or "").lower()

    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        questions = _parse_xlsx(data)
    elif name.endswith(".csv"):
        questions = _parse_csv(data)
    elif name.endswith(".txt"):
        questions = _parse_text(data.decode("utf-8-sig", errors="ignore"))
    elif name.endswith(".docx"):
        questions = _parse_text(_parse_docx(data))
    elif name.endswith(".pdf"):
        questions = _parse_text(_parse_pdf(data))
    else:
        _add_warning("Format non pris en charge (utilisez .xlsx, .csv, .docx, .pdf ou .txt).")
        questions = []

    return {"questions": questions, "warnings": list(_WARN), "nb": len(questions)}
