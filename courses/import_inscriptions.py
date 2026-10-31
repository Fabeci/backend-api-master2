# courses/import_inscriptions.py
# =============================================================================
# Extraction d'une LISTE D'APPRENANTS depuis un document, pour créer des
# inscriptions à un cours :
#   - CSV / Excel (.csv, .xlsx) : colonnes email / nom / prénom / matricule
#     (+ classe / filière facultatives pour les inscriptions établissement)
#   - Word (.docx) : paragraphes + tableaux
#   - PDF : lignes de texte
#   - TXT : une entrée par ligne
# =============================================================================

import csv
import io
import logging
import re
import unicodedata

logger = logging.getLogger(__name__)

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
WORDS_RE = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ'’\-]+")

_HEADER_KEYS = {
    "email": {"email", "e mail", "mail", "courriel", "adresse email", "adresse mail"},
    "nom": {"nom", "name", "nom de famille", "lastname", "last name", "nom et prenom", "nom prenom"},
    "prenom": {"prenom", "firstname", "first name", "prenoms"},
    "apprenant_id": {"apprenant", "apprenant id", "id", "matricule", "identifiant", "num", "numero"},
    "cours": {"cours", "course", "formation"},
    "classe": {"classe", "classe nom", "groupe", "niveau"},
    "filiere": {"filiere", "filiere nom", "filière", "filière nom"},
}

_SKIP_WORDS = {
    "liste", "inscription", "inscriptions", "apprenant", "apprenants", "etudiant", "etudiants",
    "email", "mail", "courriel", "nom", "prenom", "matricule", "cours", "classe", "groupe",
    "filiere",
}


def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def _read_csv(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig", errors="ignore")
    sample = text[:3000]
    delim = ";" if sample.count(";") > sample.count(",") else ","
    return [[(c or "").strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delim)]


def _read_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if v is None else str(v).strip() for v in row])
    return rows


def _rows_to_learners(rows: list[list[str]]) -> list[dict]:
    rows = [[str(c).strip() for c in r] for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return []

    header = [_norm(c) for c in rows[0]]
    idx = {}
    for i, h in enumerate(header):
        for key, aliases in _HEADER_KEYS.items():
            if h in aliases:
                idx.setdefault(key, i)

    has_header = any(k in idx for k in ("email", "nom", "prenom", "apprenant_id"))
    data = rows[1:] if has_header else rows
    out = []

    for r in data:
        def cell(key):
            i = idx.get(key)
            return r[i] if i is not None and i < len(r) else ""

        email = cell("email")
        m = EMAIL_RE.search(email) if email else None
        if not m:
            for c in r:
                m = EMAIL_RE.search(c)
                if m:
                    break
        email = (m.group(0) if m else "").lower()

        nom = cell("nom")
        prenom = cell("prenom")
        aid = cell("apprenant_id")
        cours = cell("cours")
        classe = cell("classe")
        filiere = cell("filiere")

        mapped = {i for i in idx.values() if i is not None}
        texts = [
            c for j, c in enumerate(r)
            if j not in mapped and c and not EMAIL_RE.search(c) and not c.strip().isdigit()
            and _norm(c) not in _SKIP_WORDS
        ]
        if not nom and not prenom and texts:
            if len(texts) >= 2:
                nom, prenom = texts[0], texts[1]
            else:
                nom = texts[0]

        if not aid:
            nums = [c for c in r if c.strip().isdigit()]
            if nums and not email and not nom:
                aid = nums[0]

        if not (email or nom or prenom or aid):
            continue

        out.append({
            "email": email,
            "nom": nom,
            "prenom": prenom,
            "apprenant_id": aid,
            "cours": cours,
            "classe": classe,
            "filiere": filiere,
            "raw": " | ".join(r),
        })
    return out


def _is_header_line(s: str) -> bool:
    n = _norm(s)
    if not n:
        return True
    if n in _SKIP_WORDS:
        return True
    if not EMAIL_RE.search(s) and (n.startswith("liste") or n.startswith("inscriptions")):
        return True
    return False


def _lines_to_learners(lines: list[str]) -> list[dict]:
    out = []
    seen = set()

    for line in lines:
        s = re.sub(r"\s+", " ", (line or "").strip())
        if not s or len(s) > 200 or _is_header_line(s):
            continue

        emails = EMAIL_RE.findall(s)
        if emails:
            rest = EMAIL_RE.sub(" ", s)
            words = [w for w in WORDS_RE.findall(rest) if _norm(w) not in _SKIP_WORDS]
            nom = " ".join(words[:2]) if words else ""
            for e in emails:
                key = e.lower()
                if key in seen:
                    continue
                seen.add(key)
                out.append({
                    "email": key, "nom": nom, "prenom": "",
                    "apprenant_id": "", "cours": "", "classe": "", "raw": s,
                })
            continue

        words = WORDS_RE.findall(s)
        if 2 <= len(words) <= 4:
            key = _norm(s)
            if key in seen:
                continue
            seen.add(key)
            out.append({
                "email": "", "nom": s, "prenom": "",
                "apprenant_id": "", "cours": "", "raw": s,
            })

    return out


def _read_docx(data: bytes):
    import docx

    d = docx.Document(io.BytesIO(data))
    lines = [(p.text or "").strip() for p in d.paragraphs if (p.text or "").strip()]
    tables = []
    for t in d.tables:
        rows = [[(c.text or "").strip() for c in row.cells] for row in t.rows]
        tables.append(rows)
    return lines, tables


def _read_pdf(data: bytes) -> list[str]:
    from courses import pdf_style

    info = pdf_style.pdf_lines(data)
    return [l["text"] for p in info["pages"] for l in p["lines"]]


SUPPORTED = (".csv", ".xlsx", ".xlsm", ".docx", ".pdf", ".txt")


def extract_learners(filename: str, data: bytes) -> dict:
    name = (filename or "").lower()
    warnings = []
    learners = []

    if name.endswith(".csv"):
        learners = _rows_to_learners(_read_csv(data))
    elif name.endswith((".xlsx", ".xlsm")):
        learners = _rows_to_learners(_read_xlsx(data))
    elif name.endswith(".docx"):
        lines, tables = _read_docx(data)
        learners = _lines_to_learners(lines)
        for rows in tables:
            learners += _rows_to_learners(rows)
    elif name.endswith(".pdf"):
        learners = _lines_to_learners(_read_pdf(data))
    elif name.endswith(".txt"):
        text = data.decode("utf-8-sig", errors="ignore")
        learners = _lines_to_learners(text.splitlines())
    else:
        raise ValueError("Format non pris en charge : CSV, Excel, Word, PDF ou TXT.")

    # Déduplication (email prioritaire, sinon nom normalisé).
    seen = set()
    unique = []
    for r in learners:
        key = (r.get("email") or "").lower() or _norm(r.get("nom") or r.get("raw") or "")
        if not key or key in seen:
            continue
        seen.add(key)
        unique.append(r)

    if not unique:
        warnings.append(
            "Aucun apprenant détecté. Formats attendus : une adresse email par ligne, "
            "ou un tableau avec les colonnes email / nom / prénom."
        )

    return {"rows": unique, "warnings": warnings}
