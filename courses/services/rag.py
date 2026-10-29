# courses/services/rag.py
"""RAG documentaire léger pour le tuteur IA.

Extrait le texte des documents du cours (PDF, DOCX, TXT, MD, CSV, XLSX,
images et PDF scannés via OCR), le découpe en extraits et les indexe dans
`DocumentChunk`. La recherche est lexicale
(pas de base vectorielle ni de dépendance lourde) et strictement limitée au
cours demandé — donc aux documents auxquels l'apprenant a déjà accès.
"""
from __future__ import annotations

import logging
import os
import re
import unicodedata
from pathlib import Path

from django.conf import settings

from courses.models import BlocContenu, DocumentChunk, RessourceSequence

logger = logging.getLogger(__name__)

TAILLE_CHUNK = 1200
CHEVAUCHEMENT = 200
MAX_PAGES = 60
MAX_CARACTERES = 250_000
MAX_CHUNKS_PAR_DOC = 300
MAX_RESSOURCES_PAR_SEQUENCE = 8

# Mots vides français (recherche lexicale).
STOPWORDS = {
    "alors", "aussi", "autre", "autres", "avec", "avoir", "beaucoup", "bien",
    "cela", "celui", "celle", "ces", "cet", "cette", "chaque", "comme",
    "comment", "dans", "depuis", "des", "deux", "donc", "dont", "elle",
    "elles", "encore", "entre", "est", "être", "faire", "fait", "fois",
    "font", "juste", "leur", "leurs", "lors", "mais", "même", "moins",
    "nous", "parce", "pour", "pourquoi", "peut", "plus", "quand", "que",
    "quel", "quelle", "quelles", "quels", "qui", "sans", "selon", "ses",
    "sont", "sous", "sur", "tout", "tous", "toute", "toutes", "tres", "très",
    "une", "vers", "voici", "votre", "vous", "avec", "explique", "expliquer",
    "donne", "donner", "exemple", "exemples", "simple", "simplement",
    "mots", "notion", "concept", "comprendre", "compris", "aide", "aider",
    "peux", "veux", "vouloir", "sais", "savoir", "dire", "parle", "parler",
}


def _nettoyer(texte: str) -> str:
    texte = (texte or "").replace("\x00", " ")
    texte = re.sub(r"[\u0000-\u0008\u000b\u000c\u000e-\u001f]", " ", texte)
    texte = re.sub(r"[ \t]+", " ", texte)
    texte = re.sub(r"\n{3,}", "\n\n", texte)
    return texte.strip()[:MAX_CARACTERES]


def _sans_accents(texte: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", texte or "")
        if unicodedata.category(c) != "Mn"
    )


def mots_cles(question: str) -> list[str]:
    """Tokens significatifs de la question (sans accents, sans mots vides)."""
    base = _sans_accents((question or "").lower())
    tokens = re.findall(r"[a-z0-9]{4,}", base)
    vus = []
    for token in tokens:
        if token in STOPWORDS or token in vus:
            continue
        vus.append(token)
        if len(vus) >= 10:
            break
    return vus


# ── Extraction ──────────────────────────────────────────────────────────────

def extraire_texte(fichier) -> list[tuple[int | None, str]]:
    """Retourne [(page|None, texte)] pour un FileField donné."""
    if not fichier:
        return []
    nom = (getattr(fichier, "name", "") or "").lower()
    try:
        chemin = fichier.path
    except Exception:  # noqa: BLE001
        return []
    if not chemin or not os.path.exists(chemin):
        return []

    try:
        if nom.endswith(".pdf"):
            import pdfplumber

            pages: list[tuple[int | None, str]] = []
            with pdfplumber.open(chemin) as pdf:
                for index, page in enumerate(pdf.pages[:MAX_PAGES], start=1):
                    texte = page.extract_text() or ""
                    if texte.strip():
                        pages.append((index, texte))
            if pages:
                return pages

            # PDF sans texte natif (scan) : OCR page par page.
            try:
                from courses import ocr
                if ocr.disponible():
                    textes = ocr.ocr_pdf_pages(
                        Path(chemin).read_bytes(), max_pages=MAX_PAGES,
                    )
                    return [
                        (index, texte)
                        for index, texte in enumerate(textes, start=1)
                        if (texte or "").strip()
                    ]
            except Exception:  # noqa: BLE001
                logger.exception("[rag] OCR PDF impossible pour %s", nom)
            return []

        if nom.endswith((".docx", ".doc")):
            import docx

            document = docx.Document(chemin)
            texte = "\n".join(p.text for p in document.paragraphs if p.text.strip())
            for table in document.tables:
                for row in table.rows:
                    texte += "\n" + " | ".join(c.text.strip() for c in row.cells)
            return [(None, texte)] if texte.strip() else []

        if nom.endswith((".txt", ".md", ".markdown", ".csv")):
            with open(chemin, "r", encoding="utf-8", errors="ignore") as flux:
                return [(None, flux.read(MAX_CARACTERES))]

        if nom.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook

            wb = load_workbook(chemin, read_only=True, data_only=True)
            morceaux: list[str] = []
            for ws in wb.worksheets:
                morceaux.append(f"Feuille : {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    valeurs = [str(v) for v in row if v is not None and str(v).strip()]
                    if valeurs:
                        morceaux.append(" | ".join(valeurs))
            return [(None, "\n".join(morceaux)[:MAX_CARACTERES])]

        if nom.endswith((".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp")):
            try:
                from courses import ocr
                if ocr.disponible():
                    texte = ocr.ocr_image(Path(chemin).read_bytes())
                    return [(1, texte)] if texte.strip() else []
            except Exception:  # noqa: BLE001
                logger.exception("[rag] OCR image impossible pour %s", nom)
            return []
    except Exception:  # noqa: BLE001
        logger.exception("[rag] extraction impossible pour %s", nom)
    return []


def decouper(texte: str, taille: int = TAILLE_CHUNK,
             chevauchement: int = CHEVAUCHEMENT) -> list[str]:
    """Découpe un texte en extraits avec chevauchement."""
    texte = _nettoyer(texte)
    if not texte:
        return []
    pas = max(200, taille - chevauchement)
    morceaux: list[str] = []
    for debut in range(0, len(texte), pas):
        morceau = texte[debut:debut + taille].strip()
        if len(morceau) >= 20:
            morceaux.append(morceau)
        if debut + taille >= len(texte):
            break
    return morceaux


# ── Indexation ──────────────────────────────────────────────────────────────

def _type_source(nom: str) -> str:
    return Path(nom or "").suffix.lstrip(".").lower() or "texte"


def _indexer(cours, nom: str, pages, bloc=None, ressource=None) -> int:
    # Les chunks de transcription vidéo/audio ont leur propre cycle de vie.
    DocumentChunk.objects.filter(bloc=bloc, ressource=ressource).exclude(
        source_type="transcript"
    ).delete()
    institution = getattr(cours, "institution", None)
    a_creer: list[DocumentChunk] = []
    ordre = 0

    for page, texte in pages:
        for morceau in decouper(texte):
            a_creer.append(DocumentChunk(
                cours=cours,
                bloc=bloc,
                ressource=ressource,
                institution=institution,
                source_nom=os.path.basename(nom or "document")[:255],
                source_type=_type_source(nom),
                page=page,
                ordre=ordre,
                texte=morceau,
            ))
            ordre += 1
            if len(a_creer) >= MAX_CHUNKS_PAR_DOC:
                break
        if len(a_creer) >= MAX_CHUNKS_PAR_DOC:
            break

    if a_creer:
        DocumentChunk.objects.bulk_create(a_creer, batch_size=200)
    return len(a_creer)


def indexer_bloc(bloc) -> int:
    if not bloc:
        return 0
    media = getattr(bloc, "fichier", None) or getattr(bloc, "image", None)
    if not media:
        return 0
    cours = bloc.sequence.module.cours
    pages = extraire_texte(media)
    if not pages:
        return 0
    return _indexer(cours, media.name, pages, bloc=bloc)


def indexer_ressource(ressource) -> int:
    if not ressource or not getattr(ressource, "fichier", None):
        return 0
    cours = ressource.sequence.module.cours
    pages = extraire_texte(ressource.fichier)
    if not pages:
        return 0
    return _indexer(cours, ressource.fichier.name, pages, ressource=ressource)


def _media_path(url: str | None) -> Path | None:
    """Résout une URL média locale (/media/...) vers un fichier du disque."""
    if not url:
        return None
    try:
        from django.conf import settings
        media_url = getattr(settings, "MEDIA_URL", "/media/") or "/media/"
        if url.startswith(media_url):
            rel = url[len(media_url):].lstrip("/")
            path = Path(settings.MEDIA_ROOT) / rel
            return path if path.exists() and path.is_file() else None
    except Exception:  # noqa: BLE001
        logger.debug("[rag] résolution média impossible", exc_info=True)
    return None


def _bloc_media_path(bloc) -> Path | None:
    """Fichier audio/vidéo local d'un bloc (upload ou URL /media/)."""
    from courses.transcription import MEDIA_EXTENSIONS

    try:
        fichier = getattr(bloc, "fichier", None)
        if fichier and Path(str(fichier.name)).suffix.lower() in MEDIA_EXTENSIONS:
            try:
                return Path(fichier.path)
            except Exception:  # noqa: BLE001
                pass
        for url in (getattr(bloc, "video_url", None), getattr(bloc, "audio_url", None)):
            path = _media_path(url)
            if path and path.suffix.lower() in MEDIA_EXTENSIONS:
                return path
    except Exception:  # noqa: BLE001
        logger.debug("[rag] média du bloc indisponible", exc_info=True)
    return None


def indexer_transcription(bloc, texte: str) -> int:
    """Indexe la transcription d'un média comme source documentaire citable."""
    if not bloc or not (texte or "").strip():
        return 0
    DocumentChunk.objects.filter(bloc=bloc, source_type="transcript").delete()
    cours = bloc.sequence.module.cours
    nom = f"{(bloc.titre or 'Média').strip()}.transcript"[:255]
    institution = getattr(cours, "institution", None)
    a_creer: list[DocumentChunk] = []
    for ordre, morceau in enumerate(decouper(texte)):
        a_creer.append(DocumentChunk(
            cours=cours,
            bloc=bloc,
            institution=institution,
            source_nom=nom,
            source_type="transcript",
            page=None,
            ordre=ordre,
            texte=morceau,
        ))
        if len(a_creer) >= MAX_CHUNKS_PAR_DOC:
            break
    if a_creer:
        DocumentChunk.objects.bulk_create(a_creer, batch_size=200)
    logger.info("[rag] transcription indexée : %s (%s chunks)", nom, len(a_creer))
    return len(a_creer)


def _assurer_transcription(bloc) -> None:
    """Transcrit à la demande un média local, une seule fois."""
    if not bloc:
        return
    if DocumentChunk.objects.filter(bloc=bloc, source_type="transcript").exists():
        return
    path = _bloc_media_path(bloc)
    if not path:
        return
    try:
        from courses import transcription
        if not transcription.disponible():
            return
        texte = transcription.transcrire(path)
        if texte.strip():
            indexer_transcription(bloc, texte)
    except Exception:  # noqa: BLE001
        logger.exception("[rag] transcription impossible pour le bloc %s", bloc.pk)


def assurer_index_bloc(bloc) -> None:
    """Indexe à la demande le fichier du bloc, ses médias et les ressources."""
    if not bloc:
        return

    media = getattr(bloc, "fichier", None) or getattr(bloc, "image", None)
    if media and not DocumentChunk.objects.filter(bloc=bloc).exclude(source_type="transcript").exists():
        try:
            indexer_bloc(bloc)
        except Exception:  # noqa: BLE001
            logger.exception("[rag] indexation du bloc %s impossible", bloc.pk)

    try:
        _assurer_transcription(bloc)
    except Exception:  # noqa: BLE001
        logger.exception("[rag] préparation transcription impossible")

    try:
        ressources = bloc.sequence.ressources_sequences.all()[:MAX_RESSOURCES_PAR_SEQUENCE]
    except Exception:  # noqa: BLE001
        return

    for ressource in ressources:
        if not getattr(ressource, "fichier", None):
            continue
        if DocumentChunk.objects.filter(ressource=ressource).exists():
            continue
        try:
            indexer_ressource(ressource)
        except Exception:  # noqa: BLE001
            logger.exception("[rag] indexation de la ressource %s impossible", ressource.pk)


def indexer_cours(cours) -> int:
    total = 0
    blocs = (
        BlocContenu.objects
        .filter(sequence__module__cours=cours)
        .exclude(fichier="")
        .exclude(fichier__isnull=True)
    )
    for bloc in blocs:
        total += indexer_bloc(bloc)
    ressources = (
        RessourceSequence.objects
        .filter(sequence__module__cours=cours)
        .exclude(fichier="")
        .exclude(fichier__isnull=True)
    )
    for ressource in ressources:
        total += indexer_ressource(ressource)
    return total


# ── Recherche ───────────────────────────────────────────────────────────────

def rechercher(cours, question: str, limite: int = 3) -> list[dict]:
    """Retourne les extraits les plus pertinents du cours (jamais d'un autre cours)."""
    mots = mots_cles(question)
    if not cours or not mots:
        return []

    resultats: list[tuple[float, DocumentChunk]] = []
    requete = (
        DocumentChunk.objects
        .filter(cours=cours)
        .only("id", "texte", "source_nom", "page", "source_type", "bloc_id", "ressource_id")
    )
    for chunk in requete.iterator(chunk_size=400):
        texte = _sans_accents((chunk.texte or "").lower())
        nom = _sans_accents((chunk.source_nom or "").lower())
        score = 0.0
        for mot in mots:
            occurrences = texte.count(mot)
            if occurrences:
                score += occurrences * (1.5 if mot in nom else 1.0)
        if score <= 0:
            continue
        # Normalisation : évite de favoriser mécaniquement les extraits longs.
        resultats.append((score / (1 + len(texte) / 800.0), chunk))

    resultats.sort(key=lambda item: -item[0])

    vus = set()
    retenus: list[dict] = []
    for score, chunk in resultats:
        cle = (chunk.source_nom, chunk.page)
        if cle in vus:
            continue
        vus.add(cle)
        retenus.append({
            "texte": chunk.texte,
            "source_nom": chunk.source_nom,
            "source_type": chunk.source_type,
            "page": chunk.page,
            "bloc_id": chunk.bloc_id,
            "ressource_id": chunk.ressource_id,
            "score": round(score, 2),
        })
        if len(retenus) >= limite:
            break
    return retenus
