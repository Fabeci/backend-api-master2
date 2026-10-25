# users/import_users.py
# =============================================================================
# Extraction d'UTILISATEURS (formateurs / responsables) depuis un document :
#   CSV / Excel : colonnes email, nom, prenom, role, telephone, password, departement
#   Word / PDF / TXT : une ligne par personne (email + nom/prénom + rôle)
# =============================================================================

import csv
import io
import logging
import re
import unicodedata

logger = logging.getLogger(__name__)

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
WORDS_RE = re.compile(r"[A-Za-zÀ-ÖØ-öø-ÿ'’\-]+")

ROLE_FORMATEUR = "Formateur"
ROLE_RESPONSABLE = "ResponsableAcademique"
ROLES_VALIDES = {ROLE_FORMATEUR, ROLE_RESPONSABLE}

_ROLE_ALIASES = {
    "formateur": ROLE_FORMATEUR, "formatrice": ROLE_FORMATEUR,
    "enseignant": ROLE_FORMATEUR, "teacher": ROLE_FORMATEUR,
    "responsable": ROLE_RESPONSABLE, "responsable academique": ROLE_RESPONSABLE,
    "responsable pédagogique": ROLE_RESPONSABLE, "responsable pedagogique": ROLE_RESPONSABLE,
    "academic manager": ROLE_RESPONSABLE,
}

_HEADER_KEYS = {
    "email": {"email", "e mail", "mail", "courriel", "adresse email"},
    "nom": {"nom", "name", "nom de famille", "lastname"},
    "prenom": {"prenom", "firstname", "first name", "prenoms"},
    "role": {"role", "fonction", "profil", "type"},
    "telephone": {"telephone", "tel", "phone", "mobile", "numero"},
    "password": {"password", "mot de passe", "motdepasse", "mdp"},
    "departement": {"departement", "department", "dept"},
}

_SKIP = {
    "liste", "utilisateurs", "formateurs", "responsables", "email", "mail",
    "nom", "prenom", "role", "fonction", "telephone", "departement",
}


def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def _role(raw, defaut=ROLE_FORMATEUR):
    n = _norm(raw)
    if not n:
        return defaut
    for key, val in _ROLE_ALIASES.items():
        if key in n:
            return val
    return ""


def _read_csv(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig", errors="ignore")
    sample = text[:3000]
    delim = ";" if sample.count(";") > sample.count(",") else ","
    return [[(c or "").strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delim)]


def _read_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    return [["" if v is None else str(v).strip() for v in row] for row in ws.iter_rows(values_only=True)]


def _rows_to_users(rows, defaut=ROLE_FORMATEUR) -> list[dict]:
    rows = [[str(c).strip() for c in r] for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return []

    header = [_norm(c) for c in rows[0]]
    idx = {}
    for i, h in enumerate(header):
        for key, aliases in _HEADER_KEYS.items():
            if h in aliases:
                idx.setdefault(key, i)

    has_header = any(k in idx for k in ("email", "nom", "prenom", "role"))
    data = rows[1:] if has_header else rows
    out = []

    for r in data:
        def cell(key):
            i = idx.get(key)
            return r[i] if i is not None and i < len(r) else ""

        email = cell("email")
        m = EMAIL_RE.search(email) if email else None
        if not m:
            for c in r:
                m = EMAIL_RE.search(c)
                if m:
                    break
        email = (m.group(0) if m else "").lower()

        nom = cell("nom")
        prenom = cell("prenom")
        mapped = {i for i in idx.values() if i is not None}
        texts = [
            c for j, c in enumerate(r)
            if j not in mapped and c and not EMAIL_RE.search(c) and _norm(c) not in _SKIP
        ]
        if not nom and not prenom and texts:
            if len(texts) >= 2:
                nom, prenom = texts[0], texts[1]
            else:
                nom = texts[0]

        out.append({
            "email": email,
            "nom": nom,
            "prenom": prenom,
            "role": _role(cell("role"), defaut),
            "telephone": cell("telephone"),
            "password": cell("password"),
            "departement": cell("departement"),
            "raw": " | ".join(r),
        })
    return out


def _lines_to_users(lines, defaut=ROLE_FORMATEUR) -> list[dict]:
    out = []
    seen = set()
    for line in lines:
        s = re.sub(r"\s+", " ", (line or "").strip())
        if not s or len(s) > 220 or _norm(s) in _SKIP:
            continue
        emails = EMAIL_RE.findall(s)
        if not emails:
            continue
        rest = EMAIL_RE.sub(" ", s)
        role = _role(s, defaut)
        words = [
            w for w in WORDS_RE.findall(rest)
            if _norm(w) not in _SKIP and _norm(w) not in _ROLE_ALIASES
            and not re.fullmatch(r"\d+", w)
        ]
        nom = words[0] if words else ""
        prenom = words[1] if len(words) > 1 else ""
        for e in emails:
            key = e.lower()
            if key in seen:
                continue
            seen.add(key)
            out.append({
                "email": key, "nom": nom, "prenom": prenom,
                "role": role, "telephone": "", "password": "",
                "departement": "", "raw": s,
            })
    return out


def _read_docx(data: bytes):
    import docx
    d = docx.Document(io.BytesIO(data))
    lines = [(p.text or "").strip() for p in d.paragraphs if (p.text or "").strip()]
    tables = [
        [[(c.text or "").strip() for c in row.cells] for row in t.rows]
        for t in d.tables
    ]
    return lines, tables


def _read_pdf(data: bytes) -> list[str]:
    from courses import pdf_style
    info = pdf_style.pdf_lines(data)
    return [l["text"] for p in info["pages"] for l in p["lines"]]


SUPPORTED = (".csv", ".xlsx", ".xlsm", ".docx", ".pdf", ".txt")


def extract_users(filename: str, data: bytes, role_defaut: str = ROLE_FORMATEUR) -> dict:
    name = (filename or "").lower()
    users = []

    if name.endswith(".csv"):
        users = _rows_to_users(_read_csv(data), role_defaut)
    elif name.endswith((".xlsx", ".xlsm")):
        users = _rows_to_users(_read_xlsx(data), role_defaut)
    elif name.endswith(".docx"):
        lines, tables = _read_docx(data)
        users = _lines_to_users(lines, role_defaut)
        for rows in tables:
            users += _rows_to_users(rows, role_defaut)
    elif name.endswith(".pdf"):
        users = _lines_to_users(_read_pdf(data), role_defaut)
    elif name.endswith(".txt"):
        text = data.decode("utf-8-sig", errors="ignore")
        users = _lines_to_users(text.splitlines(), role_defaut)
    else:
        raise ValueError("Format non pris en charge : CSV, Excel, Word, PDF ou TXT.")

    # Déduplication par email.
    seen = set()
    unique = []
    for u in users:
        key = (u.get("email") or "").lower()
        if not key or key in seen:
            continue
        seen.add(key)
        unique.append(u)

    warnings = []
    if not unique:
        warnings.append(
            "Aucun utilisateur détecté. Fournissez au minimum une colonne email "
            "et les colonnes nom / prénom / rôle."
        )

    return {"rows": unique, "warnings": warnings}


def temp_password(n: int = 12) -> str:
    import secrets
    import string
    alphabet = string.ascii_letters + string.digits + "!@#$%&*"
    return "".join(secrets.choice(alphabet) for _ in range(n))
